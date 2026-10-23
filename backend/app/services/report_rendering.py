"""Pure report presentation: serializable inputs only; no database/configuration/device imports."""

from __future__ import annotations

import html
import json
import math
from datetime import datetime, timezone
from io import BytesIO
from typing import Any, Protocol

from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.services.snapshot_data import json_object, object_value


class ReportTest(Protocol):
    test_id: str
    operator_id: str
    start_time: str | None
    end_time: str | None
    end_reason: str | None
    sample_label: str | None
    notes: str | None


class ReportAlarm(Protocol):
    alarm_code: str
    level: int
    occur_time: str
    clear_time: str | None
    text: str | None


class ReportParameters(Protocol):
    fw_version: str | None
    param_crc: str | None
    params_json: str


def now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def number(value: Any) -> float | None:
    if isinstance(value, bool):
        return None
    try:
        value = float(value)
        return value if math.isfinite(value) else None
    except (ValueError, TypeError, OverflowError):
        return None


METRIC_ROWS = (
    ("炉温峰值", "furnace_pv_max", 1, " ℃"),
    ("料层温度峰值", "burden_temp_max", 1, " ℃"),
    ("最大压差 ΔPmax", "delta_p_max", 0, " Pa"),
    ("ΔPmax 对应料层温度", "delta_p_max_temp", 1, " ℃"),
    ("总滴落量", "drip_weight_total", 2, " g"),
    ("滴落温度 Td", "td_drip_temp", 1, " ℃"),
    ("最大位移", "displacement_max", 2, " mm"),
    ("T10（10% 收缩温度）", "t10", 1, " ℃"),
    ("T40（40% 收缩温度）", "t40", 1, " ℃"),
    ("熔化开始温度 Ts", "ts", 1, " ℃"),
    ("熔落带厚度 ΔH", "delta_h_mm", 1, " mm"),
    ("T40 − T10", "t40_minus_t10", 1, " ℃"),
    ("Td − Ts", "td_minus_ts", 1, " ℃"),
    ("Td − T10", "td_minus_t10", 1, " ℃"),
    ("600 ℃ 位移基准", "reference_displacement_mm", 2, " mm"),
)


def _num(v: Any) -> float | None:
    return number(v)


def _fmt(v: Any, digits: int = 1, unit: str = "") -> str:
    n = _num(v)
    return "N/A" if n is None else f"{n:.{digits}f}{unit}"


def _xlsx_value(value: Any) -> Any:
    """阻止外部文本在电子表格中被解释为公式。"""
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _provenance_rows(test, metrics):
    basis, _ = json_object(getattr(test, "measurement_basis_json", None))
    context = object_value(basis.get("report_context"))
    specimen = object_value(basis.get("sample_metadata"))
    recipe = object_value(metrics.get("recipe_snapshot"))
    validation = object_value(recipe.get("validation"))
    rows = [
        ("参考标准", "GB/T 34211-2017"),
        ("实验室", context.get("laboratory_name") or "未记录"),
        ("实验室地址", context.get("laboratory_address") or "未记录"),
        ("试验日期", context.get("test_date") or "未记录"),
        ("异常操作", context.get("abnormal_operations") or "未记录"),
        ("标准未规定的操作", context.get("additional_operations") or "未记录"),
        ("条件记录版本", basis.get("metadata_revision", 0)),
        ("样品与装样条件", json.dumps(specimen, ensure_ascii=False, sort_keys=True) if specimen else "未记录"),
        ("配方身份/版本", f"{recipe.get('recipe_id', 'unknown')} / {recipe.get('version', 'unknown')}"),
        ("配方摘要", recipe.get("digest") or "未记录"),
        ("工艺偏离项", ", ".join(validation.get("deviations", [])) or "未记录偏离项；不能据此认定符合标准"),
        ("判定规则依据", basis.get("rules_reference") or "未确认"),
        ("固件/协议", f"{recipe.get('firmware', 'unknown')} / {recipe.get('protocol_version', 'unknown')}"),
        ("实验模式", {"standard": "标准模板", "custom": "非标"}.get(metrics.get("mode"), "未知")),
        ("算法版本", metrics.get("algorithm_version", "unknown")),
        ("符合性状态", metrics.get("compliance", "not_certified")),
        ("全实验数据完整性", metrics.get("data_integrity", "unknown")),
        ("测定段数据完整性", metrics.get("measurement_data_integrity", "unknown")),
        ("测定窗口采样点", str(metrics.get("measurement_sample_count", 0))),
        ("测定窗口外保留点", str(metrics.get("excluded_sample_count", 0))),
        ("600 ℃ 位移基准来源", metrics.get("reference_source") or "缺失"),
        ("结果限制", ", ".join(metrics.get("limitations", [])) or "无自动检出的数据缺口"),
        ("符合性", "未签发国标符合性结论；缺失数据、原文争议和现场条件须复核"),
    ]

    if "recovery" in basis:
        rows.extend(
            [
                ("档案来源", "恢复档案；缺失资料保持未知，发现时间不代表真实开始时间"),
                ("发现时间", getattr(test, "discovered_at", None) or "未知"),
            ]
        )
    profile = object_value(recipe.get("safety_profile"))
    if str(profile.get("profile_id", "")).startswith("SIMULATOR-ONLY/"):
        rows.append(("数据来源", "模拟实验；无实体执行器，未经真机认证（not_certified）"))
    return rows


def _render_html(
    test: ReportTest,
    metrics: dict[str, Any],
    sample_count: int,
    alarms: list[ReportAlarm],
    params: ReportParameters | None,
) -> str:
    e = html.escape

    def row(k: str, v: str) -> str:
        return f"<tr><th>{e(k)}</th><td>{e(v)}</td></tr>"

    h_note = "" if metrics.get("original_height_mm") else "（未提供原始料层高度 H，无法计算）"
    metric_rows = "".join(
        row(
            label,
            _fmt(metrics[key], digits, unit) + (f" {h_note}" if key in {"t10", "t40"} and metrics[key] is None else ""),
        )
        for label, key, digits, unit in METRIC_ROWS
    )

    alarm_rows = (
        "".join(
            f"<tr><td>L{a.level}</td><td>{e(a.alarm_code)}</td><td>{e(a.text or '')}</td>"
            f"<td>{e(a.occur_time)}</td><td>{e(a.clear_time or '—')}</td></tr>"
            for a in alarms[:50]
        )
        or "<tr><td colspan='5'>无报警</td></tr>"
    )

    params_summary = ""
    if params is not None:
        try:
            pv = json.loads(params.params_json)
            params_summary = f"<pre>{e(json.dumps(pv, ensure_ascii=False, indent=2))}</pre>"
        except (ValueError, TypeError):
            params_summary = "<p>参数快照解析失败</p>"
    else:
        params_summary = "<p>无本试验参数快照</p>"

    return f"""<!DOCTYPE html>
<html lang="zh-CN"><head><meta charset="UTF-8">
<title>熔滴炉试验报告 {e(test.test_id)}</title>
<style>
 body {{ font-family: 'Segoe UI','Microsoft YaHei',sans-serif; color:#1a2230; margin:32px; }}
 h1 {{ font-size:20px; border-bottom:2px solid #2563eb; padding-bottom:8px; }}
 h2 {{ font-size:15px; margin-top:24px; color:#2563eb; }}
 table {{ border-collapse:collapse; width:100%; margin-top:8px; font-size:13px; }}
 th,td {{ border:1px solid #d0d7e2; padding:6px 10px; text-align:left; }}
 th {{ background:#f1f5fb; width:220px; }}
 .meta th {{ width:160px; }}
 pre {{ background:#f7f9fc; border:1px solid #d0d7e2; padding:10px; font-size:12px; overflow:auto; }}
 .foot {{ margin-top:28px; color:#6b7686; font-size:11px; }}
</style></head><body>
<h1>熔滴炉试验报告</h1>
<table class="meta">
 {row("试验编号", test.test_id)}
 {row("操作员", test.operator_id)}
 {row("开始时间", test.start_time or "未知")}
 {row("结束时间", test.end_time or "进行中")}
 {row("结束原因", test.end_reason or "—")}
 {row("样品标识", test.sample_label or "—")}
 {row("原始料层高度 H", _fmt(metrics.get("original_height_mm"), 2, " mm"))}
 {row("备注", test.notes or "—")}
 {row("采样点数", str(sample_count))}
 {row("固件版本", (params.fw_version if params else None) or "—")}
 {row("参数 CRC", (params.param_crc if params else None) or "—")}
</table>

<h2>计算依据与结果限制</h2>
<table>{"".join(row(k, str(v)) for k, v in _provenance_rows(test, metrics))}</table>
<h2>结果指标（GB/T 34211-2017 §9）</h2>
<table>{metric_rows}</table>

<h2>报警汇总</h2>
<table><tr><th>等级</th><th>报警码</th><th>说明</th><th>发生</th><th>消除</th></tr>{alarm_rows}</table>

<h2>试验开始参数快照</h2>
{params_summary}

<p class="foot">生成时间：{metrics.get("report_generated_at") or now_iso()} · 本报告由 smd-web-hmi 自动生成 · 原始曲线与事件日志见日志导出包。</p>
</body></html>"""


def _render_pdf(
    test: ReportTest,
    metrics: dict[str, Any],
    sample_count: int,
    alarms: list[ReportAlarm],
    params: ReportParameters | None,
) -> bytes:
    """生成适合打印/归档的 Unicode PDF。"""
    buffer = BytesIO()
    font_name = "STSong-Light"
    pdfmetrics.registerFont(UnicodeCIDFont(font_name))
    styles = getSampleStyleSheet()
    for style_name in ("Title", "Heading2", "BodyText"):
        styles[style_name].fontName = font_name
    styles["Heading2"].keepWithNext = True
    document = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=16 * mm,
        rightMargin=16 * mm,
        topMargin=14 * mm,
        bottomMargin=14 * mm,
        title=f"熔滴炉试验报告 {test.test_id}",
    )
    story = [Paragraph("熔滴炉试验报告", styles["Title"]), Spacer(1, 4 * mm)]

    metadata = [
        ["试验编号", test.test_id],
        ["操作员", test.operator_id],
        ["开始时间", test.start_time or "未知"],
        ["结束时间", test.end_time or "进行中"],
        ["结束原因", test.end_reason or "—"],
        ["样品标识", test.sample_label or "—"],
        ["原始料层高度 H", _fmt(metrics.get("original_height_mm"), 2, " mm")],
        ["备注", test.notes or "—"],
        ["采样点数", str(sample_count)],
        ["固件版本", (params.fw_version if params else None) or "—"],
        ["参数 CRC", (params.param_crc if params else None) or "—"],
    ]
    table_style = TableStyle(
        [
            ("FONTNAME", (0, 0), (-1, -1), font_name),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
            ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#EFF6FF")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("LEFTPADDING", (0, 0), (-1, -1), 6),
            ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ]
    )
    metadata.extend(
        [
            [key, Paragraph(html.escape(str(value)), styles["BodyText"])]
            for key, value in _provenance_rows(test, metrics)
        ]
    )
    meta_table = Table(metadata, colWidths=[43 * mm, 120 * mm], repeatRows=0)
    meta_table.setStyle(table_style)
    story.extend([meta_table, Spacer(1, 5 * mm), Paragraph("结果指标", styles["Heading2"])])
    metric_data = [["指标", "结果"]] + [
        [label.replace("−", "-"), _fmt(metrics[key], digits, unit)] for label, key, digits, unit in METRIC_ROWS
    ]
    metrics_table = Table(metric_data, colWidths=[82 * mm, 81 * mm], repeatRows=1)
    metrics_table.setStyle(table_style)
    metrics_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DBEAFE")),
                ("FONTNAME", (0, 0), (-1, 0), font_name),
            ]
        )
    )
    story.extend([metrics_table, Spacer(1, 5 * mm), Paragraph("报警汇总", styles["Heading2"])])
    alarm_data = [["等级", "报警码", "说明", "发生", "消除"]]
    alarm_data.extend(
        [f"L{alarm.level}", alarm.alarm_code, alarm.text or "", alarm.occur_time, alarm.clear_time or "—"]
        for alarm in alarms[:50]
    )
    if len(alarm_data) == 1:
        alarm_data.append(["—", "无报警", "", "", ""])
    alarm_table = Table(alarm_data, colWidths=[12 * mm, 28 * mm, 51 * mm, 38 * mm, 38 * mm], repeatRows=1)
    alarm_table.setStyle(table_style)
    story.extend(
        [
            alarm_table,
            Spacer(1, 5 * mm),
            Paragraph(
                f"生成时间：{metrics.get('report_generated_at') or now_iso()} | smd-web-hmi 自动生成",
                styles["BodyText"],
            ),
        ]
    )
    document.build(story)
    return buffer.getvalue()


def _render_xlsx(
    test: ReportTest,
    metrics: dict[str, Any],
    sample_count: int,
    alarms: list[ReportAlarm],
    params: ReportParameters | None,
) -> bytes:
    """生成可继续统计加工的 XLSX 工作簿。"""
    workbook = Workbook()
    summary = workbook.active
    summary.title = "试验摘要"
    summary.append(["字段", "值"])
    for cell in summary[1]:
        cell.font = Font(bold=True)
    for item in (
        ("试验编号", test.test_id),
        ("操作员", test.operator_id),
        ("开始时间", test.start_time or "未知"),
        ("结束时间", test.end_time or "进行中"),
        ("结束原因", test.end_reason or "—"),
        ("样品标识", test.sample_label or "—"),
        ("原始料层高度 H (mm)", metrics.get("original_height_mm")),
        ("备注", test.notes or "—"),
        ("采样点数", sample_count),
        ("固件版本", (params.fw_version if params else None) or "—"),
        ("参数 CRC", (params.param_crc if params else None) or "—"),
        ("生成时间", metrics.get("report_generated_at") or now_iso()),
    ):
        summary.append([_xlsx_value(value) for value in item])
    for item in _provenance_rows(test, metrics):
        summary.append([_xlsx_value(value) for value in item])
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 52

    metric_sheet = workbook.create_sheet("结果指标")
    metric_sheet.append(["指标", "原始值", "显示值"])
    for cell in metric_sheet[1]:
        cell.font = Font(bold=True)
    for label, key, digits, unit in METRIC_ROWS:
        metric_sheet.append([label, metrics[key], _fmt(metrics[key], digits, unit)])
    metric_sheet.column_dimensions["A"].width = 30
    metric_sheet.column_dimensions["B"].width = 16
    metric_sheet.column_dimensions["C"].width = 20

    alarm_sheet = workbook.create_sheet("报警汇总")
    alarm_sheet.append(["等级", "报警码", "说明", "发生时间", "消除时间"])
    for cell in alarm_sheet[1]:
        cell.font = Font(bold=True)
    for alarm in alarms[:50]:
        alarm_sheet.append(
            [
                _xlsx_value(value)
                for value in (alarm.level, alarm.alarm_code, alarm.text or "", alarm.occur_time, alarm.clear_time)
            ]
        )
    for column, width in zip("ABCDE", (10, 24, 48, 28, 28), strict=True):
        alarm_sheet.column_dimensions[column].width = width

    param_sheet = workbook.create_sheet("试验开始参数")
    param_sheet.append(["参数快照 JSON"])
    param_sheet["A1"].font = Font(bold=True)
    param_sheet.append([_xlsx_value(params.params_json if params is not None else "无本试验参数快照")])
    param_sheet.column_dimensions["A"].width = 100

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
