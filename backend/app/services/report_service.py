"""报告生成：共用带质量标志、600 ℃ 基准和版本来源的国标计算器。"""

from __future__ import annotations

import asyncio
import html
import json
import uuid
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import and_, func, or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import get_settings
from app.db.models import AlarmLog, ParameterSnapshot, ReportExport, SamplePoint, TestSession
from app.db.v2_models import V2LogCursor, V2LogGap, V2SourceRecord
from app.hostcomm.protocol import now_iso
from app.hostcomm.v2_contract.types import SampleRef
from app.services.snapshot_data import json_object, object_value
from app.services.standard_metrics import MetricAccumulator, V2MetricAccumulator, number
from app.services.test_id import InvalidTestIdError, validate_test_id

SUPPORTED_FORMATS = {"html", "pdf", "xlsx"}

METRIC_ROWS = (
    ("炉温峰值", "furnace_pv_max", 1, " ℃"),
    ("料层温度峰值", "burden_temp_max", 1, " ℃"),
    ("最大压差 ΔPmax", "delta_p_max", 0, " Pa"),
    ("ΔPmax 对应料层温度", "delta_p_max_temp", 1, " ℃"),
    ("总滴落量", "drip_weight_total", 2, " g"),
    ("滴落温度 Td", "td_drip_temp", 1, " ℃"),
    ("最大位移", "displacement_max", 2, " mm"),
    ("T10（10% 收缩温度）", "t10", 1, " ℃"),
    ("T40（40% 收缩温度）", "t40", 1, " ℃"),
    ("熔化开始温度 Ts", "ts", 1, " ℃"),
    ("熔落带厚度 ΔH", "delta_h_mm", 1, " mm"),
    ("T40 − T10", "t40_minus_t10", 1, " ℃"),
    ("Td − Ts", "td_minus_ts", 1, " ℃"),
    ("Td − T10", "td_minus_t10", 1, " ℃"),
    ("600 ℃ 位移基准", "reference_displacement_mm", 2, " mm"),
)


async def _write_report_file(path: Path, content: str | bytes) -> None:
    """以排他方式写报告；取消任务前先结束线程并清理本次生成的文件。"""

    def write_exclusive() -> None:
        if isinstance(content, bytes):
            with path.open("xb") as output:
                output.write(content)
        else:
            with path.open("x", encoding="utf-8") as output:
                output.write(content)

    worker = asyncio.create_task(asyncio.to_thread(write_exclusive))
    try:
        await asyncio.shield(worker)
    except asyncio.CancelledError:
        try:
            await worker
        finally:
            path.unlink(missing_ok=True)
        raise


def _num(v: Any) -> float | None:
    return number(v)


def _fmt(v: Any, digits: int = 1, unit: str = "") -> str:
    n = _num(v)
    return "N/A" if n is None else f"{n:.{digits}f}{unit}"


def _xlsx_value(value: Any) -> Any:
    """阻止外部文本在电子表格中被解释为公式。"""
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def compute_metrics(samples: list[SamplePoint], original_height_mm: float | None) -> dict[str, Any]:
    """内存样本与数据库流使用同一算法。缺失质量和参考值保持未知。"""
    reducer = MetricAccumulator(original_height_mm)
    for sample in samples:
        reducer.add(sample)
    return reducer.finish()


def _u64_at_most(column, value):
    return or_(func.length(column) < len(value), and_(func.length(column) == len(value), column <= value))


async def _v2_measurement_integrity(session, basis):
    """Derive read-only evidence from validated cuts, never from a user-editable completion flag."""
    if not isinstance(basis, dict):
        return "unknown", None
    if basis.get("outcome") in {"invalid", "aborted"}:
        return "incomplete", None
    try:
        start, end = (SampleRef.model_validate(basis.get(key)) for key in ("measurement_start", "measurement_end"))
    except (ValueError, TypeError):
        return "unknown", None
    if start.boot_id != end.boot_id or int(start.sample_seq) > int(end.sample_seq):
        return "incomplete", None
    boundaries = []
    for reference in (start, end):
        row = await session.scalar(
            select(V2SourceRecord).where(
                V2SourceRecord.device_id == basis.get("device_id"),
                V2SourceRecord.run_id == basis.get("run_id"),
                V2SourceRecord.record_type == "sample",
                V2SourceRecord.boot_id == reference.boot_id,
                V2SourceRecord.source_seq == reference.sample_seq,
            )
        )
        if row is None or row.record_bytes is None or row.log_id is None or row.record_seq is None:
            return "unknown", None
        boundaries.append(row)
    first, last = boundaries
    if first.log_id != last.log_id or int(first.record_seq) > int(last.record_seq):
        return "incomplete", None
    evidence = {
        "device_id": first.device_id,
        "log_id": first.log_id,
        "first_record_seq": first.record_seq,
        "last_record_seq": last.record_seq,
    }
    cursor = await session.get(V2LogCursor, (first.device_id, first.log_id))
    covered = (
        cursor is not None
        and cursor.verified_from_seq is not None
        and cursor.verified_through_seq is not None
        and int(cursor.verified_from_seq) <= int(first.record_seq)
        and int(cursor.verified_through_seq) >= int(last.record_seq)
    )
    if not covered:
        # An earlier reported gap remains audit evidence, but a later proven full
        # cut of the same immutable log can fill it without deleting its history.
        gap = await session.scalar(
            select(V2LogGap.id)
            .where(
                V2LogGap.device_id == first.device_id,
                V2LogGap.log_id == first.log_id,
                _u64_at_most(V2LogGap.first_record_seq, last.record_seq),
                (
                    ~_u64_at_most(V2LogGap.last_record_seq, str(int(first.record_seq) - 1))
                    if int(first.record_seq)
                    else True
                ),
            )
            .limit(1)
        )
        return "incomplete" if gap is not None else "unknown", evidence
    evidence.update(verified_from_seq=cursor.verified_from_seq, verified_through_seq=cursor.verified_through_seq)
    return "complete" if basis.get("measurement_complete") is True else "incomplete", evidence


async def compute_metrics_from_database(
    session: AsyncSession,
    test_id: str,
    original_height_mm: float | None,
) -> dict[str, Any]:
    """分批流式读取；v2 使用原始 sample 引用，v1 保留既有接收顺序。"""
    test = await session.scalar(select(TestSession).where(TestSession.test_id == test_id))
    basis, basis_valid = json_object(getattr(test, "measurement_basis_json", None))
    boundary = basis.get("measurement_end_sample_id")
    if boundary is not None and (type(boundary) is not int or boundary <= 0):
        basis_valid, boundary = False, None
    measurement_integrity = basis.get("measurement_data_integrity", getattr(test, "data_integrity", "unknown"))
    if not isinstance(measurement_integrity, str) or measurement_integrity not in {"unknown", "complete", "incomplete"}:
        basis_valid, measurement_integrity = False, "unknown"
    recipe_raw = getattr(test, "recipe_snapshot_json", None)
    recipe, recipe_valid = json_object(recipe_raw if recipe_raw != "null" else None)
    is_v2 = "v2" in basis
    source_log_evidence = None
    if is_v2:
        measurement_integrity, source_log_evidence = await _v2_measurement_integrity(session, basis.get("v2"))
    accumulator_type = V2MetricAccumulator if is_v2 else MetricAccumulator
    v2_options = {"v2_basis": basis.get("v2"), "profile_snapshot": recipe.get("safety_profile")} if is_v2 else {}
    reducer = accumulator_type(
        original_height_mm,
        **v2_options,
        measurement_complete=basis_valid and getattr(test, "measurement_completed_at", None) is not None,
        # V2's fixed event contract is checked against every raw detector quality below.
        detector_verified=is_v2 or basis.get("detector_verified") is True,
        data_complete=basis_valid and measurement_integrity == "complete",
        measurement_end_sample_id=boundary,
        test_id=test_id,
    )
    if not basis_valid:
        reducer.limitations.add("malformed_measurement_basis")
    if not recipe_valid:
        reducer.limitations.add("malformed_recipe_snapshot")
    # Decimal uint64 strings must not be cast to SQLite's signed 64-bit INTEGER.
    # Boot identity is a grouping key; only proven same-boot boundaries enter v2 metrics.
    ordering = (
        [
            SamplePoint.source_boot_id,
            func.length(SamplePoint.source_sequence),
            SamplePoint.source_sequence,
            SamplePoint.id,
        ]
        if is_v2
        else [SamplePoint.id]
    )
    rows = await session.stream_scalars(
        select(SamplePoint).where(SamplePoint.test_id == test_id).order_by(*ordering).execution_options(yield_per=512)
    )
    try:
        async for sample in rows:
            reducer.add(sample)
    finally:
        await rows.close()
    result = reducer.finish()
    if is_v2:
        if result.pop("source_measurement_invalid"):
            measurement_integrity = "incomplete"
        elif measurement_integrity == "complete" and "freshness_profile_missing_or_invalid" in result["limitations"]:
            measurement_integrity = "unknown"
        result["measurement_source_log"] = source_log_evidence
    result.update(
        standard="GB/T 34211-2017",
        mode=getattr(test, "mode", "custom"),
        data_integrity=getattr(test, "data_integrity", "unknown"),
        measurement_data_integrity=measurement_integrity,
        recipe_snapshot=recipe if recipe else None,
        compliance="not_certified",
    )
    return result


def _provenance_rows(test, metrics):
    basis, _ = json_object(getattr(test, "measurement_basis_json", None))
    context = object_value(basis.get("report_context"))
    specimen = object_value(basis.get("sample_metadata"))
    recipe = object_value(metrics.get("recipe_snapshot"))
    validation = object_value(recipe.get("validation"))
    rows = [
        ("参考标准", "GB/T 34211-2017"),
        ("实验室", context.get("laboratory_name") or "未记录"),
        ("实验室地址", context.get("laboratory_address") or "未记录"),
        ("试验日期", context.get("test_date") or "未记录"),
        ("异常操作", context.get("abnormal_operations") or "未记录"),
        ("标准未规定的操作", context.get("additional_operations") or "未记录"),
        ("条件记录版本", basis.get("metadata_revision", 0)),
        ("样品与装样条件", json.dumps(specimen, ensure_ascii=False, sort_keys=True) if specimen else "未记录"),
        ("配方身份/版本", f"{recipe.get('recipe_id', 'unknown')} / {recipe.get('version', 'unknown')}"),
        ("配方摘要", recipe.get("digest") or "未记录"),
        ("工艺偏离项", ", ".join(validation.get("deviations", [])) or "未记录偏离项；不能据此认定符合标准"),
        ("判定规则依据", basis.get("rules_reference") or "未确认"),
        ("固件/协议", f"{recipe.get('firmware', 'unknown')} / {recipe.get('protocol_version', 'unknown')}"),
        ("实验模式", "标准模板" if metrics.get("mode") == "standard" else "非标 / 历史未标定"),
        ("算法版本", metrics.get("algorithm_version", "unknown")),
        ("全实验数据完整性", metrics.get("data_integrity", "unknown")),
        ("测定段数据完整性", metrics.get("measurement_data_integrity", "unknown")),
        ("测定窗口采样点", str(metrics.get("measurement_sample_count", 0))),
        ("测定窗口外保留点", str(metrics.get("excluded_sample_count", 0))),
        ("600 ℃ 位移基准来源", metrics.get("reference_source") or "缺失"),
        ("结果限制", ", ".join(metrics.get("limitations", [])) or "无自动检出的数据缺口"),
        ("符合性", "未签发国标符合性结论；缺失数据、原文争议和现场条件须复核"),
    ]

    return rows


def _render_html(
    test: TestSession,
    metrics: dict[str, Any],
    sample_count: int,
    alarms: list[AlarmLog],
    params: ParameterSnapshot | None,
) -> str:
    e = html.escape

    def row(k: str, v: str) -> str:
        return f"<tr><th>{e(k)}</th><td>{e(v)}</td></tr>"

    h_note = "" if metrics.get("original_height_mm") else "（未提供原始料层高度 H，无法计算）"
    metric_rows = "".join(
        row(
            label,
            _fmt(metrics[key], digits, unit) + (f" {h_note}" if key in {"t10", "t40"} and metrics[key] is None else ""),
        )
        for label, key, digits, unit in METRIC_ROWS
    )

    alarm_rows = (
        "".join(
            f"<tr><td>L{a.level}</td><td>{e(a.alarm_code)}</td><td>{e(a.text or '')}</td>"
            f"<td>{e(a.occur_time)}</td><td>{e(a.clear_time or '—')}</td></tr>"
            for a in alarms[:50]
        )
        or "<tr><td colspan='5'>无报警</td></tr>"
    )

    params_summary = ""
    if params is not None:
        try:
            pv = json.loads(params.params_json)
            params_summary = f"<pre>{e(json.dumps(pv, ensure_ascii=False, indent=2))}</pre>"
        except (ValueError, TypeError):
            params_summary = "<p>参数快照解析失败</p>"
    else:
        params_summary = "<p>无本试验参数快照</p>"

    return f"""<!DOCTYPE html>
<html lang="zh-CN"><head><meta charset="UTF-8">
<title>熔滴炉试验报告 {e(test.test_id)}</title>
<style>
 body {{ font-family: 'Segoe UI','Microsoft YaHei',sans-serif; color:#1a2230; margin:32px; }}
 h1 {{ font-size:20px; border-bottom:2px solid #2563eb; padding-bottom:8px; }}
 h2 {{ font-size:15px; margin-top:24px; color:#2563eb; }}
 table {{ border-collapse:collapse; width:100%; margin-top:8px; font-size:13px; }}
 th,td {{ border:1px solid #d0d7e2; padding:6px 10px; text-align:left; }}
 th {{ background:#f1f5fb; width:220px; }}
 .meta th {{ width:160px; }}
 pre {{ background:#f7f9fc; border:1px solid #d0d7e2; padding:10px; font-size:12px; overflow:auto; }}
 .foot {{ margin-top:28px; color:#6b7686; font-size:11px; }}
</style></head><body>
<h1>熔滴炉试验报告</h1>
<table class="meta">
 {row("试验编号", test.test_id)}
 {row("操作员", test.operator_id)}
 {row("开始时间", test.start_time)}
 {row("结束时间", test.end_time or "进行中")}
 {row("结束原因", test.end_reason or "—")}
 {row("样品标识", test.sample_label or "—")}
 {row("原始料层高度 H", _fmt(metrics.get("original_height_mm"), 2, " mm"))}
 {row("备注", test.notes or "—")}
 {row("采样点数", str(sample_count))}
 {row("固件版本", (params.fw_version if params else None) or "—")}
 {row("参数 CRC", (params.param_crc if params else None) or "—")}
</table>

<h2>计算依据与结果限制</h2>
<table>{"".join(row(k, str(v)) for k, v in _provenance_rows(test, metrics))}</table>
<h2>结果指标（GB/T 34211-2017 §9）</h2>
<table>{metric_rows}</table>

<h2>报警汇总</h2>
<table><tr><th>等级</th><th>报警码</th><th>说明</th><th>发生</th><th>消除</th></tr>{alarm_rows}</table>

<h2>试验开始参数快照</h2>
{params_summary}

<p class="foot">生成时间：{now_iso()} · 本报告由 smd-web-hmi 自动生成 · 原始曲线与事件日志见日志导出包。</p>
</body></html>"""


def _render_pdf(
    test: TestSession,
    metrics: dict[str, Any],
    sample_count: int,
    alarms: list[AlarmLog],
    params: ParameterSnapshot | None,
) -> bytes:
    """生成适合打印/归档的 Unicode PDF。"""
    buffer = BytesIO()
    font_name = "STSong-Light"
    pdfmetrics.registerFont(UnicodeCIDFont(font_name))
    styles = getSampleStyleSheet()
    for style_name in ("Title", "Heading2", "BodyText"):
        styles[style_name].fontName = font_name
    document = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=16 * mm,
        rightMargin=16 * mm,
        topMargin=14 * mm,
        bottomMargin=14 * mm,
        title=f"熔滴炉试验报告 {test.test_id}",
    )
    story = [Paragraph("熔滴炉试验报告", styles["Title"]), Spacer(1, 4 * mm)]

    metadata = [
        ["试验编号", test.test_id],
        ["操作员", test.operator_id],
        ["开始时间", test.start_time],
        ["结束时间", test.end_time or "进行中"],
        ["结束原因", test.end_reason or "—"],
        ["样品标识", test.sample_label or "—"],
        ["原始料层高度 H", _fmt(metrics.get("original_height_mm"), 2, " mm")],
        ["备注", test.notes or "—"],
        ["采样点数", str(sample_count)],
        ["固件版本", (params.fw_version if params else None) or "—"],
        ["参数 CRC", (params.param_crc if params else None) or "—"],
    ]
    table_style = TableStyle(
        [
            ("FONTNAME", (0, 0), (-1, -1), font_name),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
            ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#EFF6FF")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("LEFTPADDING", (0, 0), (-1, -1), 6),
            ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ]
    )
    metadata.extend(
        [
            [key, Paragraph(html.escape(str(value)), styles["BodyText"])]
            for key, value in _provenance_rows(test, metrics)
        ]
    )
    meta_table = Table(metadata, colWidths=[43 * mm, 120 * mm], repeatRows=0)
    meta_table.setStyle(table_style)
    story.extend([meta_table, Spacer(1, 5 * mm), Paragraph("结果指标", styles["Heading2"])])
    metric_data = [["指标", "结果"]] + [
        [label, _fmt(metrics[key], digits, unit)] for label, key, digits, unit in METRIC_ROWS
    ]
    metrics_table = Table(metric_data, colWidths=[82 * mm, 81 * mm], repeatRows=1)
    metrics_table.setStyle(table_style)
    metrics_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DBEAFE")),
                ("FONTNAME", (0, 0), (-1, 0), font_name),
            ]
        )
    )
    story.extend([metrics_table, Spacer(1, 5 * mm), Paragraph("报警汇总", styles["Heading2"])])
    alarm_data = [["等级", "报警码", "说明", "发生", "消除"]]
    alarm_data.extend(
        [f"L{alarm.level}", alarm.alarm_code, alarm.text or "", alarm.occur_time, alarm.clear_time or "—"]
        for alarm in alarms[:50]
    )
    if len(alarm_data) == 1:
        alarm_data.append(["—", "无报警", "", "", ""])
    alarm_table = Table(alarm_data, colWidths=[12 * mm, 28 * mm, 51 * mm, 38 * mm, 38 * mm], repeatRows=1)
    alarm_table.setStyle(table_style)
    story.extend(
        [
            alarm_table,
            Spacer(1, 5 * mm),
            Paragraph(f"生成时间：{now_iso()} · smd-web-hmi 自动生成", styles["BodyText"]),
        ]
    )
    document.build(story)
    return buffer.getvalue()


def _render_xlsx(
    test: TestSession,
    metrics: dict[str, Any],
    sample_count: int,
    alarms: list[AlarmLog],
    params: ParameterSnapshot | None,
) -> bytes:
    """生成可继续统计加工的 XLSX 工作簿。"""
    workbook = Workbook()
    summary = workbook.active
    summary.title = "试验摘要"
    summary.append(["字段", "值"])
    for cell in summary[1]:
        cell.font = Font(bold=True)
    for item in (
        ("试验编号", test.test_id),
        ("操作员", test.operator_id),
        ("开始时间", test.start_time),
        ("结束时间", test.end_time or "进行中"),
        ("结束原因", test.end_reason or "—"),
        ("样品标识", test.sample_label or "—"),
        ("原始料层高度 H (mm)", metrics.get("original_height_mm")),
        ("备注", test.notes or "—"),
        ("采样点数", sample_count),
        ("固件版本", (params.fw_version if params else None) or "—"),
        ("参数 CRC", (params.param_crc if params else None) or "—"),
        ("生成时间", now_iso()),
    ):
        summary.append([_xlsx_value(value) for value in item])
    for item in _provenance_rows(test, metrics):
        summary.append([_xlsx_value(value) for value in item])
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 52

    metric_sheet = workbook.create_sheet("结果指标")
    metric_sheet.append(["指标", "原始值", "显示值"])
    for cell in metric_sheet[1]:
        cell.font = Font(bold=True)
    for label, key, digits, unit in METRIC_ROWS:
        metric_sheet.append([label, metrics[key], _fmt(metrics[key], digits, unit)])
    metric_sheet.column_dimensions["A"].width = 30
    metric_sheet.column_dimensions["B"].width = 16
    metric_sheet.column_dimensions["C"].width = 20

    alarm_sheet = workbook.create_sheet("报警汇总")
    alarm_sheet.append(["等级", "报警码", "说明", "发生时间", "消除时间"])
    for cell in alarm_sheet[1]:
        cell.font = Font(bold=True)
    for alarm in alarms[:50]:
        alarm_sheet.append(
            [
                _xlsx_value(value)
                for value in (alarm.level, alarm.alarm_code, alarm.text or "", alarm.occur_time, alarm.clear_time)
            ]
        )
    for column, width in zip("ABCDE", (10, 24, 48, 28, 28), strict=True):
        alarm_sheet.column_dimensions[column].width = width

    param_sheet = workbook.create_sheet("试验开始参数")
    param_sheet.append(["参数快照 JSON"])
    param_sheet["A1"].font = Font(bold=True)
    param_sheet.append([_xlsx_value(params.params_json if params is not None else "无本试验参数快照")])
    param_sheet.column_dimensions["A"].width = 100

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


async def generate_report(
    session: AsyncSession,
    test_id: str,
    *,
    operator_id: str,
    fmt: str = "html",
    options: dict[str, Any] | None = None,
) -> ReportExport:
    """生成并登记 HTML、PDF 或 XLSX 报告。"""
    test_id = validate_test_id(test_id)
    options = options or {}
    if fmt not in SUPPORTED_FORMATS:
        raise ValueError(f"不支持的报告格式: {fmt}")

    test = await session.scalar(select(TestSession).where(TestSession.test_id == test_id))
    if test is None:
        raise ValueError(f"试验不存在: {test_id}")

    sample_count = await session.scalar(
        select(func.count()).select_from(SamplePoint).where(SamplePoint.test_id == test_id)
    )
    alarms = list(
        (
            await session.execute(
                select(AlarmLog)
                .where(AlarmLog.test_id == test_id)
                .order_by(AlarmLog.level.desc(), AlarmLog.id.desc())
                .limit(50)
            )
        ).scalars()
    )
    params = await session.scalar(
        select(ParameterSnapshot)
        .where(ParameterSnapshot.test_id == test_id, ParameterSnapshot.source == "test_start")
        .order_by(ParameterSnapshot.id.desc())
    )

    original_height_mm = (
        test.original_height_mm if test.original_height_mm is not None else _num(options.get("original_height_mm"))
    )
    metrics = await compute_metrics_from_database(session, test_id, original_height_mm)
    renderers = {"html": _render_html, "pdf": _render_pdf, "xlsx": _render_xlsx}
    content = renderers[fmt](test, metrics, int(sample_count or 0), alarms, params)

    settings = get_settings()
    stamp = datetime.now().strftime("%Y%m%d%H%M%S")
    reports_dir = settings.reports_dir.resolve()
    path = (reports_dir / f"{test_id}-{stamp}-{uuid.uuid4().hex}.{fmt}").resolve()
    if not path.is_relative_to(reports_dir):
        # test_id 已有白名单；这里保留最终写入点的纵深防御。
        raise InvalidTestIdError("报告路径超出报告目录")
    await _write_report_file(path, content)
    size = (await asyncio.to_thread(path.stat)).st_size

    record = ReportExport(
        test_id=test_id,
        generated_at=now_iso(),
        operator_id=operator_id,
        format=fmt,
        file_path=str(path),
        file_size_bytes=size,
        notes=json.dumps({"metrics": metrics}, ensure_ascii=False),
    )
    session.add(record)
    # 试验会话登记报告路径
    test.report_path = str(path)
    await session.commit()
    await session.refresh(record)
    return record
