"""Recovered archives must not invent experiment metadata in any report format."""

import zipfile
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.db.models import TestSession
from app.services import report_service


def recovered_test():
    test = TestSession(test_id="RECOVERED-1", operator_id="device-recovery", start_time=None, mode="unknown")
    test.discovered_at = "2026-09-07T01:02:03Z"
    test.measurement_basis_json = '{"recovery":{"case_id":1}}'
    return test


def test_recovery_html_labels_unknown_start_separately_from_discovery():
    test = recovered_test()
    metrics = report_service.compute_metrics([], None)
    metrics["mode"] = "unknown"
    content = report_service._render_html(test, metrics, 0, [], None)
    assert "<th>开始时间</th><td>未知</td>" in content
    assert "<th>发现时间</th><td>2026-09-07T01:02:03Z</td>" in content
    assert "<th>实验模式</th><td>未知</td>" in content
    assert "恢复档案" in content


def test_recovery_xlsx_preserves_unknown_start_and_mode():
    metrics = report_service.compute_metrics([], None)
    metrics["mode"] = "unknown"
    payload = report_service._render_xlsx(recovered_test(), metrics, 0, [], None)
    workbook = load_workbook(BytesIO(payload))
    summary = dict(workbook["试验摘要"].iter_rows(min_row=2, values_only=True))
    assert summary["开始时间"] == "未知"
    assert summary["发现时间"] == "2026-09-07T01:02:03Z"
    assert summary["实验模式"] == "未知"


def test_recovery_pdf_can_render_without_actual_start():
    metrics = report_service.compute_metrics([], None)
    assert report_service._render_pdf(recovered_test(), metrics, 0, [], None).startswith(b"%PDF")


async def seed_recovery(db, *, replay_status="pending"):
    from app.db.v2_models import V2RunRecovery

    test = recovered_test()
    db.add(test)
    case = V2RunRecovery(
        id="a" * 32,
        device_id="b" * 32,
        run_id="c" * 32,
        first_seen_at=test.discovered_at,
        last_seen_at=test.discovered_at,
        evidence_json="{}",
        test_id=test.test_id,
        review_state="bound",
        replay_status=replay_status,
    )
    db.add(case)
    await db.commit()
    return test, case


@pytest.mark.parametrize("replay_status", ["pending", "running", "failed"])
async def test_recovery_report_refuses_unfinished_replay(db_session, tmp_path, monkeypatch, replay_status):
    from app.core.config import get_settings
    from app.services.v2_run_recovery import V2RecoveryError

    monkeypatch.setattr(type(get_settings()), "reports_dir", property(lambda _self: tmp_path))
    test, _ = await seed_recovery(db_session, replay_status=replay_status)
    with pytest.raises(V2RecoveryError, match="recovery_replay_pending"):
        await report_service.generate_report(db_session, test.test_id, operator_id="admin")
    assert not list(tmp_path.glob("*.html"))


async def test_recovery_report_api_rejects_before_enqueuing(db_session, monkeypatch):
    from fastapi import HTTPException

    from app.api.deps import CurrentUser
    from app.api.routes import reports

    test, _ = await seed_recovery(db_session)
    jobs = []
    monkeypatch.setattr(reports.background_jobs, "submit", lambda *args: jobs.append(args) or "d" * 32)
    with pytest.raises(HTTPException) as caught:
        await reports.generate_report(
            reports.GenerateReportRequest(test_id=test.test_id), CurrentUser("admin", "admin"), db_session
        )
    assert caught.value.status_code == 409
    assert not jobs


async def test_recovered_global_alarm_is_in_report_without_rewriting_original(db_session, tmp_path, monkeypatch):
    from app.core.config import get_settings
    from app.db.models import AlarmLog
    from app.db.v2_models import V2RecoveryAssociation

    monkeypatch.setattr(type(get_settings()), "reports_dir", property(lambda _self: tmp_path))
    test, case = await seed_recovery(db_session, replay_status="complete")
    alarm = AlarmLog(
        test_id=None, alarm_code="RECOVERY-ALARM", level=2, occur_time=test.discovered_at, text="Retained global alarm"
    )
    db_session.add(alarm)
    await db_session.flush()
    db_session.add(V2RecoveryAssociation(recovery_id=case.id, source_record_id=1, alarm_log_id=alarm.id))
    await db_session.commit()
    report = await report_service.generate_report(db_session, test.test_id, operator_id="admin")
    assert "RECOVERY-ALARM" in Path(report.file_path).read_text()
    await db_session.refresh(alarm)
    assert alarm.test_id is None

    from app.services.log_export_service import export_test_logs

    monkeypatch.setattr(type(get_settings()), "exports_dir", property(lambda _self: tmp_path))
    result = await export_test_logs(db_session, test.test_id)
    with zipfile.ZipFile(result["file_path"]) as archive:
        content = archive.read(f"{test.test_id}_alarm_log.csv").decode("utf-8")
    assert content.count("RECOVERY-ALARM") == 1


def test_synthetic_profile_is_explicitly_labeled_in_report_provenance():
    metrics = report_service.compute_metrics([], None)
    metrics["recipe_snapshot"] = {"safety_profile": {"profile_id": "SIMULATOR-ONLY/1"}}
    rows = dict(report_service._provenance_rows(recovered_test(), metrics))
    assert "模拟实验" in rows["数据来源"]
    assert "not_certified" in rows["数据来源"]


async def test_replay_finishing_during_report_does_not_publish_mixed_snapshot(db_session, tmp_path, monkeypatch):
    from sqlalchemy import update
    from sqlalchemy.ext.asyncio import async_sessionmaker

    from app.core.config import get_settings
    from app.db.models import ReportExport
    from app.db.v2_models import V2RunRecovery
    from app.services.v2_run_recovery import V2RecoveryError

    monkeypatch.setattr(type(get_settings()), "reports_dir", property(lambda _self: tmp_path))
    test, case = await seed_recovery(db_session, replay_status="complete")
    original = report_service._write_report_file
    factory = async_sessionmaker(db_session.bind, expire_on_commit=False)

    async def concurrent_replay(path, content):
        await original(path, content)
        async with factory() as other:
            await other.execute(
                update(V2RunRecovery)
                .where(V2RunRecovery.id == case.id)
                .values(
                    review_revision=V2RunRecovery.review_revision + 1,
                    replay_status="complete",
                    replay_through_id=42,
                )
            )
            await other.commit()

    monkeypatch.setattr(report_service, "_write_report_file", concurrent_replay)
    with pytest.raises(V2RecoveryError, match="recovery_changed_during_report"):
        await report_service.generate_report(db_session, test.test_id, operator_id="admin")
    assert not list(tmp_path.glob("*.html"))
    from sqlalchemy import select

    assert await db_session.scalar(select(ReportExport)) is None
